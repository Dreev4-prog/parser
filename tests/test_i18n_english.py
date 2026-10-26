"""English user journeys, literal listing data and recipient-specific delivery."""
import asyncio
import csv
import html
import re
import shutil
from datetime import datetime, timedelta
from types import SimpleNamespace as NS
from unittest.mock import AsyncMock

import pytest
from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.ext.asyncio import create_async_engine

import bot as ui
import db
import i18n
import promo_posts as promo
from categories import CATEGORIES, GROUPS
from i18n import protect_text, strip_protected, translate_text
from models import Listing, RadarProduct, UserSettings
from test_promo_admin import RecordingSession
from test_promo_posts import database


def assert_english(value):
    assert not re.search(r"[А-Яа-яЁё]", value), value
    assert not re.search(r"[\ue100\ue101\ue200\ue201]", value), value


def labels(markup):
    return "\n".join(button.text for row in markup.inline_keyboard for button in row)


def test_copy_is_natural_and_literal_data_and_html_attributes_are_preserved():
    assert translate_text("🔥 Лучшие сейчас", "en") == "🔥 Best right now"
    assert translate_text("только что", "en") == "just now"
    assert translate_text("1 день · 3 дня · 7 дней", "en") == "1 day · 3 days · 7 days"
    raw = "новые, только что, Autos, <товар>"
    url = "https://example.test/новые?tag=до&amp;sort=новые"
    source = f'🔎 Ключевые слова: <b>{protect_text(html.escape(raw))}</b> · <a href="{url}">Открыть</a>'
    result = translate_text(source, "en")
    assert html.escape(raw) in result and f'href="{url}"' in result
    assert ">Open</a>" in result
    assert translate_text(source, "ru") == strip_protected(source)
    assert translate_text("<code>новые</code>", "en") == "<code>новые</code>"


@pytest.mark.parametrize("raw,expected", [
    ("up to 100", "custom::100"), ("UP TO 100 EUR", "custom::100"),
    ("from 500", "custom:500:"), ("at least 500 €", "custom:500:"),
    ("from 120 to 250", "custom:120:250"), ("120-250", "custom:120:250"),
    ("до 100", "custom::100"), ("от 500", "custom:500:"),
    ("500+", "custom:500:"), ("-200", "custom::200"), ("250-120", None),
])
def test_price_input_supports_both_languages(raw, expected):
    assert ui.parse_scan_price_input(raw) == expected


def test_home_onboarding_filters_and_navigation_are_complete(monkeypatch):
    monkeypatch.setattr(ui, "allowed", lambda uid: uid == 1)
    token = ui._UI_LANGUAGE.set("en")
    try:
        screens = [ui.home_text(2), ui.home_text(2, access_active=False),
                   ui.home_text(1, access_active=False, trial_remaining=2)]
        screens += [ui.onboarding_text(step) for step in (1, 2, 3)]
        for mode in ui.MODE_LABELS:
            settings = UserSettings(output_mode=mode, sort_mode="newest", min_views=50,
                                    smart_dedupe=True, clean_noise=False, include_words="macbook", exclude_words="defekt")
            screens.append(ui.settings_text(settings))
            screens.append(labels(ui.settings_keyboard(settings)))
        keyboards = [ui.main_keyboard(1), ui.main_keyboard(0, access_active=False, trial_remaining=2),
                     ui.radar_home_keyboard(1), ui.radar_home_keyboard(2), ui.radar_best_keyboard(1),
                     ui.radar_best_keyboard(2), ui.radar_price_keyboard(), ui.scan_price_keyboard(),
                     ui.page_limit_keyboard(), ui.page_limit_keyboard(trial=True), ui.scan_date_keyboard(),
                     ui.post_scan_keyboard(1), ui.scan_detail_keyboard(1), ui.growth_period_keyboard(1),
                     ui.stopped_job_keyboard(), ui.failed_job_keyboard(1), ui.min_views_keyboard(50),
                     ui.groups_keyboard(set())]
        keyboards += [ui.category_keyboard(key, set()) for key in GROUPS]
        screens += [labels(ui._localize_inline_markup(kb, "en")) for kb in keyboards]
        screens += [ui.radar_daily_digest_text({"hot":2,"rising":4,"radar_total":100,"signals_today":5,
                    "new_listings":20,"listings_seen":200,"categories_processed":3,"best_score_today":90}, paid=paid)
                    for paid in (True, False)]
        for screen in screens:
            assert_english(translate_text(screen, "en"))
        assert "Kitchen &amp; Dining" in translate_text("Küche &amp; Esszimmer", "en")
        for cat in CATEGORIES.values():
            if not cat.is_group:
                assert cat.name in i18n.EN_CATEGORY_NAMES
    finally:
        ui._UI_LANGUAGE.reset(token)


def test_scan_progress_and_subscription_variants(monkeypatch):
    async def scenario():
        job = ui.ScanJob(job_id="english", user_id=1, chat_id=1, status_message_id=10,
                         category_keys=["hg_kueche"], created_at=datetime.utcnow(), target_date="2026-10-25")
        monkeypatch.setattr(ui, "category_live_progress", {})
        for state in ("queued", "running"):
            job.state = state
            for phase in ("seeking", "collecting", "views"):
                job.current_progress_key = "en"
                job.current_category = "Küche & Esszimmer"
                ui.category_live_progress["en"] = ui.CategoryLiveProgress(
                    "hg_kueche", job.current_category, "newest", phase=phase, today_seen=12, views_ready=8)
                assert_english(translate_text(ui.render_user_job_status(job), "en"))
        monkeypatch.setattr(ui, "ADMIN_IDS", set())
        monkeypatch.setattr(ui, "get_commerce_user", AsyncMock(return_value=NS(access_until=datetime.utcnow()+timedelta(days=3))))
        monkeypatch.setattr(ui, "user_payments", AsyncMock(return_value=[NS(status="pending")]))
        monkeypatch.setattr(ui, "get_trial_status", AsyncMock(return_value=NS(eligible=True, remaining=2)))
        monkeypatch.setattr(ui, "providers_status", lambda:{"cryptobot":True,"xrocket":True})
        monkeypatch.setattr(ui, "is_banned_cached", lambda uid:False)
        monkeypatch.setattr(ui, "allowed", lambda uid:False)
        monkeypatch.setattr(ui, "get_plans", AsyncMock(return_value=[NS(title=t,price_usdt=5,key=str(i))
                            for i,t in enumerate(("1 день", "3 дня", "7 дней", "30 дней"))]))
        for mode in ("subscription", "open", "admin_only"):
            monkeypatch.setattr(ui, "current_access_mode", lambda:mode)
            for until in (None, datetime.utcnow()+timedelta(days=3)):
                ui.get_commerce_user.return_value = NS(access_until=until)
                assert_english(translate_text(await ui.subscription_text(1), "en"))
                assert_english(labels(ui._localize_inline_markup(await ui.subscription_keyboard(1), "en")))
    asyncio.run(scenario())


def test_radar_feeds_cards_search_and_referral_preserve_content(monkeypatch):
    async def scenario():
        now = datetime.utcnow()
        product = RadarProduct(id=1,title="новые товары & Autos",category_key="hg_kueche",status="rising",
            current_score=78,peak_score=90,confidence=72,min_price_eur=50,max_price_eur=120,best_views=177,
            best_views_per_hour=33,signal_count=3,listing_count=2,first_radar_at=now,last_signal_at=now,
            demand_views=40,latest_reason="Unified 48H: Strong · DT Score 78/100",
            opportunity_type="observed_demand",latest_source="radar3_observed")
        monkeypatch.setattr(ui, "allowed", lambda uid:uid==1)
        monkeypatch.setattr(ui, "radar_stats", AsyncMock(return_value=NS(total=65,hot=0,rising=1,fast_sold=5)))
        monkeypatch.setattr(ui, "get_radar_product", AsyncMock(return_value=(product,Listing(url="https://example.test/новые"),[])))
        monkeypatch.setattr(ui, "get_fast_sold_info", AsyncMock(return_value=None))
        monkeypatch.setattr(ui, "get_fast_sold_infos", AsyncMock(return_value={}))
        monkeypatch.setattr(ui, "is_radar_favorite", AsyncMock(return_value=False))
        monkeypatch.setattr(ui, "get_trial_status", AsyncMock(return_value=NS(eligible=True,remaining=2)))
        monkeypatch.setattr(ui, "list_radar_products", AsyncMock(return_value=([product],15)))
        monkeypatch.setattr(ui, "search_radar_products", AsyncMock(return_value=([product],1)))
        for uid in (1,2):
            home = translate_text(await ui._radar_home_text(uid), "en")
            assert_english(home)
            assert "<b>65</b>" in home and "last 7 days" in home
        for status in ui.RADAR_STATUS_LABEL:
            product.status = status
            card, kb = await ui._radar_product_payload(1,1)
            translated = translate_text(card,"en")
            assert "новые товары &amp; Autos" in translated
            assert "Kitchen &amp; Dining" in translated
            assert_english(translated.replace("новые товары", "products"))
        for mode in ("hot","rising","alltime","favorites","fastsold","category_best"):
            text_value, kb = await ui._radar_list_payload(1,mode,0,"hg_kueche" if mode=="category_best" else None)
            assert_english(translate_text(text_value,"en").replace("новые товары","products"))
            localized = ui._localize_inline_markup(kb,"en")
            assert "новые товары & Autos" in labels(localized)
        text_value, _ = await ui._radar_search_payload("новые")
        translated = translate_text(text_value,"en")
        assert "Search term: <b>новые</b>" in translated
        ui.search_radar_products.assert_awaited_once_with("новые",page=0,price_filter="any")
        monkeypatch.setattr(ui,"referral_promo_enabled",AsyncMock(return_value=True))
        monkeypatch.setattr(ui,"referral_user_stats",AsyncMock(return_value={"total":3,"eligible":3,"progress":1,"days_earned":1}))
        monkeypatch.setattr(ui,"get_user_language",AsyncMock(return_value="en"))
        text_value, kb = await ui._referral_screen(1,NS(get_me=AsyncMock(return_value=NS(username="example_bot"))))
        assert_english(translate_text(text_value,"en"))
        assert "Try%20DT%20Parser" in kb.inline_keyboard[0][0].url
    asyncio.run(scenario())


def test_background_delivery_uses_recipient_language_and_keeps_admin_context(monkeypatch):
    async def scenario():
        monkeypatch.setattr(ui,"get_user_language",AsyncMock(side_effect=lambda uid:"en" if uid==2 else "ru"))
        session=RecordingSession(); client=ui.LocalizedBot("123:TEST",session=session)
        token=ui._UI_LANGUAGE.set("ru"); owner=ui._UI_LANGUAGE_USER.set(1)
        try:
            await client.send_message(2,"🔥 Лучшие сейчас")
            await client.send_message(1,"🔥 Лучшие сейчас")
            assert session.requests[-2].text == "🔥 Best right now"
            assert session.requests[-1].text == "🔥 Лучшие сейчас"
        finally:
            ui._UI_LANGUAGE.reset(token); ui._UI_LANGUAGE_USER.reset(owner)
        await client.session.close()
    asyncio.run(scenario())


def test_exports_translate_headers_and_categories_without_changing_listings():
    row=Listing(external_id="english-export",title="новые Autos",category_key="hg_kueche",category="Küche & Esszimmer",
        price_eur=42,view_count=100,posted_date_msk="2026-10-25",posted_text="Heute",url="https://example.test/новые")
    path=ui.write_listing_xlsx([row],"all",language="en")
    try:
        ws=load_workbook(path).active
        assert ws.title == "Results" and ws["C2"].value == "Title"
        assert ws["B3"].value == "Kitchen & Dining" and ws["C3"].value == row.title
        assert ws["H3"].hyperlink.target == row.url and ws["D3"].number_format == "0"
        assert_english(" ".join(str(c.value) for c in ws[2]))
    finally:
        shutil.rmtree(path.parent)
    token=ui._UI_LANGUAGE.set("en")
    try:
        path=ui.write_listing_csv([row],"all")
        with path.open(encoding="utf-8-sig",newline="") as stream:
            values=list(csv.reader(stream,delimiter=";"))
        assert_english(" ".join(values[0]))
        assert "новые Autos" in values[1] and row.url in values[1]
    finally:
        ui._UI_LANGUAGE.reset(token); shutil.rmtree(path.parent)


def test_promo_english_delivery_and_photo_cache_are_separate(monkeypatch):
    async def scenario():
        async with database(monkeypatch):
            await promo.initialize_promotions()
            await i18n.set_user_language(1,"ru"); await i18n.set_user_language(2,"en")
            posts,_=await promo.list_posts(); payload=promo.post_payload(posts[0])
            session=RecordingSession(); client=ui.LocalizedBot("123:TEST",session=session)
            await promo.send_post(client,1,payload)
            await promo.send_post(client,2,payload)
            assert session.requests[0].photo.path.name == "promo_parser_alf.png"
            assert session.requests[1].photo.path.name == "promo_parser_alf_en.png"
            assert session.requests[1].caption == promo.PARSER_CAPTION_EN
            assert_english(labels(session.requests[1].reply_markup))
            assert len(payload["_photo_cache"]) == 2
            await promo.send_post(client,2,payload)
            assert isinstance(session.requests[-1].photo,str)
            await client.session.close()
    asyncio.run(scenario())


def test_existing_database_gets_additive_english_promo_columns(monkeypatch,tmp_path):
    async def scenario():
        engine=create_async_engine(f"sqlite+aiosqlite:///{tmp_path / 'upgrade.db'}")
        monkeypatch.setattr(db,"engine",engine)
        monkeypatch.setattr(db,"_IS_POSTGRES",False)
        async with engine.begin() as connection:
            await connection.execute(text("CREATE TABLE promo_posts (id INTEGER PRIMARY KEY, title TEXT)"))
            await connection.execute(text("INSERT INTO promo_posts (title) VALUES ('Existing post')"))
        try:
            await db.init_db()
            await db.init_db()
            async with engine.connect() as connection:
                row=(await connection.execute(text("SELECT title, caption_en, photo_en FROM promo_posts"))).one()
                assert tuple(row) == ("Existing post","","")
        finally:
            await engine.dispose()
    asyncio.run(scenario())


def test_home_uses_separate_english_and_russian_menu_images(monkeypatch):
    async def scenario():
        monkeypatch.setattr(ui,"get_selected",AsyncMock(return_value=set()))
        monkeypatch.setattr(ui,"get_settings",AsyncMock(return_value=NS(auto_observations=False)))
        monkeypatch.setattr(ui,"get_trial_status",AsyncMock(return_value=NS(eligible=False,remaining=0)))
        monkeypatch.setattr(ui,"referral_promo_enabled",AsyncMock(return_value=True))
        monkeypatch.setattr(ui,"allowed",lambda uid:True)
        monkeypatch.setattr(ui,"get_user_language",AsyncMock(side_effect=lambda uid:"en" if uid==2 else "ru"))
        monkeypatch.setattr(ui,"_MENU_IMAGE_FILE_IDS",{})
        session=RecordingSession(); client=ui.LocalizedBot("123:TEST",session=session)
        for uid in (1,2,2,1):
            message=NS(answer_photo=lambda **kwargs:client.send_photo(uid,**kwargs))
            await ui._send_home_message(message,uid)
        assert session.requests[0].photo.path.name == "dt_parser_menu.png"
        assert session.requests[1].photo.path.name == "dt_parser_menu_en.png"
        assert_english(session.requests[1].caption)
        assert all(isinstance(r.photo,str) for r in session.requests[2:])
        assert set(ui._MENU_IMAGE_FILE_IDS) == {"ru","en"}
        await client.session.close()
    asyncio.run(scenario())


@pytest.mark.parametrize("writer", [ui.write_frequent_xlsx, ui.write_market_xlsx,
    ui.write_disappearing_xlsx, ui.write_price_drop_xlsx])
def test_all_specialized_export_headers_are_in_english(writer):
    path=writer([],language="en")
    try:
        ws=load_workbook(path).active
        assert_english(str(ws["A1"].value))
        assert_english(" ".join(str(c.value) for c in ws[2]))
        assert ws.freeze_panes == "A3" and ws.auto_filter.ref.startswith("A2:")
    finally:
        shutil.rmtree(path.parent)
